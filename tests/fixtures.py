"""Genera documentos inventados para las pruebas (nada con derechos de autor ni datos reales)."""

import io
import zipfile

import pymupdf


def png(ancho=120, alto=80):
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, ancho, alto), False)
    pix.set_rect(pix.irect, (200, 30, 30))
    return pix.tobytes("png")


CONTAINER = """<?xml version="1.0"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>
</container>"""

CAPITULO = """<?xml version="1.0" encoding="utf-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>x</title></head><body>
<p class="titulo" id="cap1">Capítulo uno</p>
<p>La <b>fisiología </b>del riñón explica la filtración glomerular.^[[1]]</p>
<div>Primer bloque</div><div>Segundo bloque</div>
<p>Más datos en josephsoninstitute.org sobre ética.</p>
<p>El cartel decía: <i>No</i> <i>hay devoluciones</i>.</p>
<img src="../Images/figura%201.png" alt="Nefrona"/>
<table><tr><th>Fármaco</th><th>Dosis</th></tr><tr><td>Amoxicilina</td><td>500 mg</td></tr></table>
<ol><li>Primero</li><li>Segundo</li></ol>
<ul>
        <li>Uno</li>
        <li>Dos</li>
</ul>
<p class="titulo" id="cap2">CAPÍTULO DOS</p>
<p>Texto del segundo capítulo con eñes y tildes: corazón, señal.</p>
</body></html>"""


def epub(con_ncx=True):
    manifest = [
        '<item id="c1" href="Text/cap%201.xhtml" media-type="application/xhtml+xml"/>',
        '<item id="img" href="Images/figura%201.png" media-type="image/png"/>',
    ]
    if con_ncx:
        manifest.append('<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>')
    else:
        manifest.append('<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>')
    opf = f"""<?xml version="1.0"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Libro de prueba</dc:title><dc:creator>Autora Inventada</dc:creator></metadata>
<manifest>{''.join(manifest)}</manifest>
<spine><itemref idref="c1"/></spine></package>"""
    ncx = """<?xml version="1.0"?><ncx xmlns="http://www.daisy.org/z3986/2005/ncx/"><navMap>
<navPoint><navLabel><text>Capítulo uno</text></navLabel><content src="Text/cap%201.xhtml#cap1"/></navPoint>
<navPoint><navLabel><text>Capítulo dos</text></navLabel><content src="Text/cap%201.xhtml#cap2"/></navPoint>
</navMap></ncx>"""
    nav = """<html xmlns="http://www.w3.org/1999/xhtml"><body><nav><ol>
<li><a href="Text/cap%201.xhtml#cap1">Capítulo uno</a></li><li><a href="Text/cap%201.xhtml#cap2">Capítulo dos</a></li>
</ol></nav></body></html>"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", "application/epub+zip")
        z.writestr("META-INF/container.xml", CONTAINER)
        z.writestr("OEBPS/content.opf", opf)
        z.writestr("OEBPS/Text/cap 1.xhtml", CAPITULO)
        z.writestr("OEBPS/Images/figura 1.png", png())
        z.writestr("OEBPS/toc.ncx" if con_ncx else "OEBPS/nav.xhtml", ncx if con_ncx else nav)
    return buf.getvalue()


PARRAFO = ("La hipertensión arterial es un factor de riesgo cardiovascular modificable. "
           "El tratamiento combina cambios en el estilo de vida y fármacos antihipertensivos. ")


def pdf(paginas=3, escaneado=False, dos_columnas=False):
    doc = pymupdf.open()
    for i in range(paginas):
        pag = doc.new_page()
        if escaneado:
            pag.insert_image(pymupdf.Rect(50, 50, 550, 750), stream=png(400, 600))
            continue
        pag.insert_text((72, 80), f"Sección {i + 1}", fontsize=20)
        if dos_columnas:
            pag.insert_textbox(pymupdf.Rect(50, 110, 290, 760), PARRAFO * 4, fontsize=10)
            pag.insert_textbox(pymupdf.Rect(310, 110, 550, 760), PARRAFO * 4, fontsize=10)
        else:
            titulo = "Guía inventada de hipertensión. " if i == 0 else ""
            pag.insert_textbox(pymupdf.Rect(72, 110, 540, 760), titulo + PARRAFO * 6, fontsize=11)
    if not escaneado:
        doc[0].insert_image(pymupdf.Rect(72, 600, 272, 740), stream=png(300, 200))
        doc[0].insert_image(pymupdf.Rect(300, 600, 310, 610), stream=png(10, 10))  # viñeta: debe descartarse
    doc.set_metadata({"title": "Guía inventada de hipertensión", "author": "Equipo de prueba"})
    datos = doc.tobytes()
    doc.close()
    return datos


def docx(ruta):
    from docx import Document
    from docx.shared import Inches

    d = Document()
    d.add_heading("Protocolo inventado", 1)
    d.add_heading("Objetivos", 2)
    d.add_paragraph("Describir el manejo inicial de la sepsis en adultos.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Variable", "Valor"
    t.cell(1, 0).text, t.cell(1, 1).text = "Lactato", "2 mmol/L"
    d.add_picture(io.BytesIO(png()), width=Inches(1))
    d.save(ruta)


def pptx(ruta):
    from pptx import Presentation

    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Clase de farmacología"
    s.placeholders[1].text = "Betabloqueantes: indicaciones y contraindicaciones"
    p.save(ruta)


def xlsx(ruta):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["Paciente ficticio", "Edad"])
    wb.active.append(["Caso A", 54])
    wb.save(ruta)
