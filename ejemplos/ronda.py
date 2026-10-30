"""Prepara una ronda de revisión: convierte el corpus y arma un lote de encargos para agy (Antigravity).

Uso: uv run python ejemplos/ronda.py <carpeta-del-lote>

Cada subcarpeta del lote trae los .md generados, el texto del original extraído por otra vía (Calibre para
EPUB, texto plano de PyMuPDF para PDF, python-pptx/openpyxl para Office), los informes JSON, la checklist y un
encargo.md. agy no ejecuta nada: solo lee esos archivos y responde con hallazgos en JSON.
"""

import json
import os
import shutil
import subprocess
import sys

AQUI = os.path.dirname(os.path.abspath(__file__))
CORPUS = os.path.join(AQUI, "corpus")
SALIDAS = os.path.join(AQUI, "salidas")
CALIBRE = shutil.which("ebook-convert") or r"C:\Program Files\Calibre2\ebook-convert.exe"

GRUPOS = {
    "epub-cuatro-mil": ["Cuatro mil semanas - Gestión del tiempo para mortales (Oliver Burkeman).epub"],
    "epub-carnegie": ["Cómo ganar amigos e influir sobre las personas (CARNEGIE, DALE).epub"],
    "pdf-plos-tablas": ["plos-medicine.pdf"],
    "pdf-dos-columnas": ["arxiv-2610.34749.pdf", "arxiv-2610.38948.pdf"],
    "pdf-espanol-ocr-office": ["upch-rmh.pdf", "upch-escaneado.pdf", "protocolo-inventado.docx",
                               "clase-inventada.pptx", "datos-inventados.xlsx"],
}
CON_OCR = {"upch-escaneado.pdf"}

ENCARGO = """# Encargo

## Contexto
Revisas la calidad de una CLI que convierte documentos a Markdown para Obsidian y LLM (todo-a-markdown).
En esta carpeta están, por cada documento: `<nombre>.md` (lo que generó la CLI), `<nombre>.original.txt`
(el texto del original extraído por OTRA herramienta, como referencia; puede traer saltos de línea,
encabezados y pies de página que el .md quita a propósito) y `informes.jsonl` (el informe de calidad que
emitió la CLI, una línea por documento). `checklist.md` define los criterios A1–C7.
{nota}

## Objetivo
Encontrar los defectos reales de los .md según la checklist (secciones A, B y C), con evidencia.

## Archivos
{archivos}

## Restricciones
- **Terminal: PROHIBIDA.** No ejecutes ningún comando (ni python, node, ls, cat ni git): una sola petición
  deja VACÍA tu respuesta y se pierde todo el trabajo. Todo lo que necesitas ya está en los archivos de la lista.
- **Web: PROHIBIDA.** No abras URL ni busques en internet.
- Lee solo los archivos de la lista; no busques otros. No edites ni crees archivos.
- Compara al menos 3 tramos de cada documento largo (inicio, mitad y final) contra su .original.txt.
- Un hallazgo sin evidencia textual exacta no vale. No reportes como defecto lo que la CLI quita a propósito:
  encabezados y pies de página repetidos, números de página, saltos de línea del PDF, guiones de corte.
- No reportes problemas que ya están en el original (por ejemplo, un símbolo que en .original.txt también
  sale mal), salvo que el informe de la CLI no los avise (criterio C1/C2).

## Entregable
Tu respuesta final ES el entregable: solo el JSON del esquema, sin preámbulo. Como máximo 8 hallazgos, los
más graves primero. `linea` es el número de línea en el .md (aproximado si no puedes contarlo). En
`evidencia_md` y `evidencia_original` copia el texto exacto (máximo 200 caracteres cada uno).
`criterios_sin_problemas` lista los IDs que revisaste y están bien.

## Terminado cuando
Revisaste cada documento de la lista contra su original y su informe.
"""

ESQUEMA = {
    "type": "object",
    "required": ["hallazgos", "criterios_sin_problemas", "resumen"],
    "properties": {
        "hallazgos": {"type": "array", "maxItems": 8, "items": {
            "type": "object",
            "required": ["criterio", "gravedad", "archivo", "linea", "descripcion", "evidencia_md"],
            "properties": {
                "criterio": {"type": "string"},
                "gravedad": {"type": "string", "enum": ["critico", "mayor", "menor"]},
                "archivo": {"type": "string"},
                "linea": {"type": "integer"},
                "descripcion": {"type": "string"},
                "evidencia_md": {"type": "string"},
                "evidencia_original": {"type": "string"},
            }}},
        "criterios_sin_problemas": {"type": "array", "items": {"type": "string"}},
        "resumen": {"type": "string"},
    },
}


def texto_original(ruta, destino):
    ext = os.path.splitext(ruta)[1].lower()
    if ext == ".epub":
        subprocess.run([CALIBRE, ruta, destino], check=True, capture_output=True, stdin=subprocess.DEVNULL)
        return
    if ext == ".pdf":
        import pymupdf
        with pymupdf.open(ruta) as doc:
            texto = "".join(f"\n=== página {i + 1} ===\n{p.get_text()}" for i, p in enumerate(doc))
    elif ext == ".docx":
        from docx import Document
        d = Document(ruta)
        texto = "\n".join(p.text for p in d.paragraphs)
        texto += "\n" + "\n".join(" | ".join(c.text for c in fila.cells) for t in d.tables for fila in t.rows)
    elif ext == ".pptx":
        from pptx import Presentation
        texto = "\n".join(s.text_frame.text for d in Presentation(ruta).slides for s in d.shapes if s.has_text_frame)
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        texto = "\n".join(" | ".join(str(c) for c in fila) for fila in load_workbook(ruta).active.iter_rows(values_only=True))
    with open(destino, "w", encoding="utf-8") as f:
        f.write(texto)


def convertir(archivos, extra):
    proc = subprocess.run(["todo-a-markdown", *archivos, "-o", SALIDAS, "--sobrescribir", *extra],
                          capture_output=True, stdin=subprocess.DEVNULL)
    return [json.loads(ln) for ln in proc.stdout.decode("utf-8").splitlines() if ln.strip()]


def main(lote):
    os.makedirs(SALIDAS, exist_ok=True)
    todos = [n for g in GRUPOS.values() for n in g]
    informes = convertir([os.path.join(CORPUS, n) for n in todos if n not in CON_OCR], [])
    informes += convertir([os.path.join(CORPUS, n) for n in CON_OCR], ["--ocr"])
    por_nombre = {os.path.basename(r["archivo"]): r for r in informes}

    for grupo, nombres in GRUPOS.items():
        carpeta = os.path.join(lote, grupo)
        os.makedirs(carpeta, exist_ok=True)
        shutil.copy(os.path.join(AQUI, "checklist.md"), carpeta)
        lista, notas = ["- checklist.md", "- informes.jsonl"], []
        with open(os.path.join(carpeta, "informes.jsonl"), "w", encoding="utf-8") as f:
            for n in nombres:
                r = dict(por_nombre[n])
                base = os.path.splitext(n)[0]
                r["archivo"], r["salida"] = n, base + ".md"  # sin rutas locales del usuario
                f.write(json.dumps(r, ensure_ascii=False) + "\n")
                if r["estado"] == "error":
                    notas.append(f"- {n}: la CLI dio error ({r.get('error')}).")
                    continue
                shutil.copy(os.path.join(SALIDAS, base + ".md"), carpeta)
                texto_original(os.path.join(CORPUS, n), os.path.join(carpeta, base + ".original.txt"))
                lista += [f"- {base}.md", f"- {base}.original.txt"]
                adj = os.path.join(SALIDAS, base + "_adjuntos")
                if os.path.isdir(adj):
                    with open(os.path.join(carpeta, base + ".adjuntos.txt"), "w", encoding="utf-8") as fa:
                        fa.write("\n".join(sorted(os.listdir(adj))))
                    lista.append(f"- {base}.adjuntos.txt (lista de archivos que existen en {base}_adjuntos/)")
        if "upch-escaneado.pdf" in nombres:
            notas.append("- upch-escaneado.pdf es una versión escaneada (imagen) de upch-rmh.pdf, convertida con --ocr:"
                         " su original de referencia es el texto de upch-rmh.original.txt (solo las 2 primeras páginas).")
        with open(os.path.join(carpeta, "encargo.md"), "w", encoding="utf-8") as f:
            f.write(ENCARGO.format(archivos="\n".join(lista), nota="\n".join(notas)))
        with open(os.path.join(carpeta, "esquema.json"), "w", encoding="utf-8") as f:
            json.dump(ESQUEMA, f, ensure_ascii=False, indent=1)
    for r in informes:
        print(r["estado"], os.path.basename(r["archivo"]), r.get("error", ""), r.get("motivos", ""))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main(sys.argv[1])
